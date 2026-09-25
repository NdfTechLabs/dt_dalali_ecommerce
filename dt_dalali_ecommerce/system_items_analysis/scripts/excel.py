from openpyxl import load_workbook


def select_from_list(items, title):
    print(f"\n{title}")
    print("-" * len(title))

    for index, item in enumerate(items, start=1):
        print(f"{index}. {item}")

    while True:
        try:
            choice = int(input("\nEnter your choice: "))

            if 1 <= choice <= len(items):
                return items[choice - 1]

            print(f"Enter a number between 1 and {len(items)}.")

        except ValueError:
            print("Please enter a valid number.")


def load_excel():
    file_path = input("Enter the Excel workbook path: ").strip()

    try:
        workbook = load_workbook(file_path, read_only=True)

        print(f"\nWorkbook loaded: {file_path}")

        # Select worksheet
        worksheet_name = select_from_list(
            workbook.sheetnames,
            "Available Worksheets"
        )

        worksheet = workbook[worksheet_name]

        # Get column names from first row
        columns = [
            cell.value
            for cell in worksheet[1]
            if cell.value is not None
        ]

        if not columns:
            raise ValueError("The worksheet does not contain column headers.")

        # Select search column
        selected_column = select_from_list(
            columns,
            "Available Columns"
        )

        return workbook, worksheet, selected_column

    except FileNotFoundError:
        print("File not found.")
        raise

    except Exception as error:
        print(f"Could not open workbook: {error}")
        raise