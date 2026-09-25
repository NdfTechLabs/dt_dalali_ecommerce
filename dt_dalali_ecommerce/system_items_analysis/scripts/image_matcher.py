from pathlib import Path
import shutil
import re

from openpyxl import load_workbook

from excel import select_from_list


# Destination inside your project
DESTINATION = Path("/home/franklyne/work/Dhana/technologies/sites/dhanatechnologies.local/public/files")


IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".gif",
    ".avif"
}


def normalize_name(value):
    """
    Normalize a product/file name so that small differences
    in spaces, hyphens, underscores and capitalization don't
    prevent a match.
    """

    if value is None:
        return ""

    value = str(value).lower().strip()

    # Remove file extension only if it is a known image extension
    image_extensions = {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".gif",
        ".avif",
    }

    for extension in image_extensions:
        if value.endswith(extension):
            value = value[:-len(extension)]
            break

    # Replace anything that isn't a letter/number with a space
    value = re.sub(r"[^a-z0-9]+", " ", value)

    # Remove duplicate whitespace
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def load_image_files(images_path):
    """
    Load image files and create a lookup based on normalized
    filenames.
    """

    images_path = Path(images_path)

    if not images_path.exists():
        raise FileNotFoundError(
            f"Image directory does not exist: {images_path}"
        )

    if not images_path.is_dir():
        raise ValueError(
            f"Not a directory: {images_path}"
        )

    image_files = {}

    for file in images_path.iterdir():

        if not file.is_file():
            continue

        if file.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        normalized = normalize_name(file.name)

        image_files[normalized] = file

    return image_files


def find_column_index(worksheet, column_name):
    """
    Find the Excel column index based on the header.
    """

    for cell in worksheet[1]:

        if cell.value == column_name:
            return cell.column

    raise ValueError(
        f"Column '{column_name}' was not found."
    )


def main():

    # ---------------------------------------------
    # Workbook
    # ---------------------------------------------

    workbook_path = input(
        "Enter the Excel workbook path: "
    ).strip()

    workbook = load_workbook(workbook_path)

    # ---------------------------------------------
    # Select worksheet
    # ---------------------------------------------

    worksheet_name = select_from_list(
        workbook.sheetnames,
        "Available Worksheets"
    )

    worksheet = workbook[worksheet_name]

    print(
        f"\nSelected worksheet: {worksheet_name}"
    )

    # ---------------------------------------------
    # Get columns
    # ---------------------------------------------

    columns = [
        cell.value
        for cell in worksheet[1]
        if cell.value is not None
    ]

    if not columns:
        raise ValueError(
            "The worksheet does not contain columns."
        )

    # ---------------------------------------------
    # Select product/name column
    # ---------------------------------------------

    product_column = select_from_list(
        columns,
        "Select Product Name Column"
    )

    # ---------------------------------------------
    # Select image column
    # ---------------------------------------------

    image_column = select_from_list(
        columns,
        "Select Image Column"
    )

    # ---------------------------------------------
    # Images directory
    # ---------------------------------------------

    images_path = input(
        "\nEnter the downloaded images directory: "
    ).strip()

    # ---------------------------------------------
    # Load images
    # ---------------------------------------------

    image_files = load_image_files(images_path)

    print(
        f"\nFound {len(image_files)} image files."
    )

    # ---------------------------------------------
    # Destination
    # ---------------------------------------------

    DESTINATION.mkdir(
        parents=True,
        exist_ok=True
    )

    print(
        f"Destination: {DESTINATION.resolve()}"
    )

    # ---------------------------------------------
    # Find Excel columns
    # ---------------------------------------------

    product_column_index = find_column_index(
        worksheet,
        product_column
    )

    image_column_index = find_column_index(
        worksheet,
        image_column
    )

    # ---------------------------------------------
    # Process products
    # ---------------------------------------------

    matched = 0
    unmatched = 0

    for row in range(2, worksheet.max_row + 1):

        product_name = worksheet.cell(
            row=row,
            column=product_column_index
        ).value

        if not product_name:
            continue

        normalized_product = normalize_name(
            product_name
        )

        image_file = image_files.get(
            normalized_product
        )

        if image_file is None:

           print(
               f"[NOT FOUND] {product_name}"
           )

           print(
               f"Normalized: {normalized_product}"
           )

           unmatched += 1
           continue

        # -----------------------------------------
        # Copy image
        # -----------------------------------------

        destination_file = (
            DESTINATION / image_file.name
        )

        shutil.copy2(
            image_file,
            destination_file
        )

        # -----------------------------------------
        # Set image field
        # -----------------------------------------

        image_url = (
            f"/files/media/product/{image_file.name}"
        )

        worksheet.cell(
            row=row,
            column=image_column_index
        ).value = image_url

        matched += 1

        print(
            f"[MATCHED] {product_name}"
        )

        print(
            f"          {image_file.name}"
        )

        print(
            f"          {image_url}"
        )

    # ---------------------------------------------
    # Save workbook
    # ---------------------------------------------

    output_path = Path(workbook_path).with_name(
        f"{Path(workbook_path).stem}_updated"
        f"{Path(workbook_path).suffix}"
    )

    workbook.save(output_path)

    # ---------------------------------------------
    # Summary
    # ---------------------------------------------

    print("\n" + "=" * 60)
    print("IMAGE MATCHING COMPLETE")
    print("=" * 60)

    print(f"Matched:   {matched}")
    print(f"Unmatched: {unmatched}")

    print(
        f"\nUpdated workbook:\n"
        f"{output_path}"
    )


if __name__ == "__main__":
    main()